import json
import mimetypes
import os
import base64
from idlelib.iomenu import encoding
import requests
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from html import escape
from pathlib import Path

from gmail_tools.utils import get_all_user_emails, get_plico_logo
from utils import get_user_profile
from gmail_tools.admin_connection_local import get_gmail_service


TEST_ROOT = Path("temp_files") / ""
LOGO_FILE_GDID = "1pIsq8tpweJChbQNOsK0JAjvWE_ycviMC&usp=drive_fs"
SIGNATURE_DOC = TEST_ROOT / "EMAIL_SIGNATURE.docx"

PLICO_LOGO = TEST_ROOT / "PLICO DESIGN STUDIO_CURRENT.png"
TEMP_OUT_HTML = TEST_ROOT / "OUT"


def replace_placeholders(user_data: dict, src_html):
    out_html = src_html
    for key, value in user_data.items():
        print(key)
        out_html = out_html.replace(f"__{key}__", value)
    print(f"replaced:{out_html}")
    return out_html


def docx_to_styled_html(docx_path=SIGNATURE_DOC, logo_path=PLICO_LOGO):
    document = Document(str(docx_path))
    html_parts = [
        '<!DOCTYPE html>',
        '<html>',
        '<head><meta charset="UTF-8"></head>',
        '<body>'
    ]

    # Track inline shapes for image matching
    shape_index = 0
    inline_shapes = document.inline_shapes
    for para in document.paragraphs:
        line = ""
        for run in para.runs:
            # Check if this run contains an inline shape (image)
            if shape_index < len(inline_shapes) and run._element.xpath(".//pic:pic"):
                shape = inline_shapes[shape_index]
                shape_index += 1

                # Get size
                width_px = int(shape.width.pt * 1.333)  # convert pt to px
                height_px = int(shape.height.pt * 1.333)

                logo_link = get_plico_logo()
                if logo_link:
                    img_html = (
                        f'<img src="{logo_link}" '
                        f'width="{width_px}" height="{height_px}" alt="Logo" />'
                    )
                    line += img_html
                    continue
                else:
                    print(f"⚠️ Logo file not found: {logo_link}")

            # Process text with styles
            text = escape(run.text)
            style = []

            if run.bold:
                style.append("font-weight: bold;")
            if run.italic:
                style.append("font-style: italic;")
            if run.underline:
                style.append("text-decoration: underline;")
            if run.font.name:
                style.append(f"font-family: '{run.font.name}';")
            if run.font.size:
                pt = run.font.size.pt
                style.append(f"font-size: {pt:.1f}pt;")

            if style:
                text = f'<span style="{" ".join(style)}">{text}</span>'
            line += text

        if line.strip():
            html_parts.append(f"<p>{line}</p>")

    html_parts.append('</body></html>')

    # with open(output_html_path, 'w', encoding='utf-8') as f:
    #     f.write('\n'.join(html_parts))

    # print(f"✅ Styled HTML saved to: {output_html_path}")
    html_out = ""
    for part in html_parts:
        # print(part)
        if not html_out:
            html_out += part
        else:
            html_out += f'\n{part}'
    # print(f"html_template:{html_out}")
    return html_out


def get_user_signature(user_email):
    """
    Retrieves the Gmail signature for the specified user.

    Args:
        user_email (str): The email of the user whose signature is being retrieved.

    Returns:
        str: The HTML signature, or None if not set.
    """
    try:
        service = get_gmail_service(user_email)

        # The user can have multiple "sendAs" aliases — usually one is "primary"
        sendas_list = service.users().settings().sendAs().list(userId='me').execute()

        for sendas in sendas_list.get('sendAs', []):
            if sendas.get('isPrimary'):
                signature = sendas.get('signature', '')
                print(f"✅ Signature found for {user_email}: {signature}...")
                return signature

        print(f"⚠️ No primary sendAs address found for {user_email}.")
        return None

    except Exception as e:
        print(f"❌ Error retrieving signature for {user_email}: {e}")
        return None


def set_signatures():
    """
    Update the Gmail signature for all users.

    """
    html_template = docx_to_styled_html()
    users_email = get_all_user_emails()
    for user_email in users_email:
        # if user_email == 'david@plicodesignstudio.com':# remove this bit when in production
        print(f"doing {user_email}")
        try:
            service = get_gmail_service(user_email)
            user_data = get_user_profile(user_email)
            new_signature_html = replace_placeholders(user_data, html_template)
            print(f"new_signature:{new_signature_html}")
            # Fetch list of "sendAs" identities
            sendas_list = service.users().settings().sendAs().list(userId=user_email).execute()

            # Find the primary sendAs entry (should match the user_email)
            sendas_entry = next(
                (entry for entry in sendas_list.get('sendAs', []) if entry['sendAsEmail'] == user_email),
                None
            )

            if not sendas_entry:
                print(f"⚠️ No sendAs identity found for {user_email}")
                return

            sendas_id = sendas_entry['sendAsEmail']

            # Perform the update
            service.users().settings().sendAs().patch(
                userId=user_email,
                sendAsEmail=sendas_id,
                body={'signature': new_signature_html}
            ).execute()

            print(f"✅ Signature updated for {user_email}")
            with open(TEMP_OUT_HTML / f"{user_data['name']}.html", 'w', encoding="utf-8") as f:
                f.write(new_signature_html)
        except Exception as e:
            print(f"❌ Error updating signature for {user_email}: {e}")


def set_outlook_signature(access_token, html_signature):
    url = "https://graph.microsoft.com/v1.0/me/mailboxSettings"
    headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json"
    }
    data = {
        "signature": html_signature
    }

    response = requests.patch(url, headers=headers, json=data)

    if response.ok:
        print("✅ Outlook signature updated.")
    else:
        print(f"❌ Failed to update signature: {response.status_code} - {response.text}")

set_signatures()

